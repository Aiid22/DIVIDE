from __future__ import annotations

import csv
import gettext
import gzip
import hashlib
import io
import json
import os
import re
import sqlite3
import tarfile
import tempfile
import time
import zipfile
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

import yaml
from defusedxml import ElementTree

from divide.config import AppConfig
from divide.models import CarrierObject, CarrierRecord, ScanWarning
from divide.recovery.encoding import infer_encoding

from .registry import HandlerRegistry, default_registry


def reconstruct_ocr_line(words: list[tuple[int, str, float, dict[str, int]]]) -> tuple[str, tuple[int, int, int, int]]:
    """Restore word order and layout-driven spacing from OCR geometry."""
    ordered = sorted(words, key=lambda item: item[0])
    pieces: list[str] = []
    previous_right: int | None = None
    median_height = sorted(item[3]["height"] for item in ordered)[len(ordered) // 2]
    for _, word, _, box in ordered:
        gap = box["left"] - previous_right if previous_right is not None else 0
        if pieces and gap > max(2, median_height * .35):
            pieces.append(" ")
        pieces.append(word)
        previous_right = box["left"] + box["width"]
    left = min(item[3]["left"] for item in ordered)
    top = min(item[3]["top"] for item in ordered)
    right = max(item[3]["left"] + item[3]["width"] for item in ordered)
    bottom = max(item[3]["top"] + item[3]["height"] for item in ordered)
    return "".join(pieces), (left, top, right - left, bottom - top)


@dataclass(slots=True)
class LocalizationResult:
    records: list[CarrierRecord] = field(default_factory=list)
    objects: list[CarrierObject] = field(default_factory=list)
    warnings: list[ScanWarning] = field(default_factory=list)
    scanned_objects: int = 0


@dataclass(slots=True)
class _State:
    result: LocalizationResult = field(default_factory=LocalizationResult)
    total_expanded: int = 0
    record_counter: int = 0
    ocr_pages: int = 0
    deadline: float | None = None
    object_counter: int = 0
    object_metadata: dict[str, dict[str, Any]] = field(default_factory=dict)
    disabled_carrier_types: frozenset[str] = frozenset()
    objects: dict[str, CarrierObject] = field(default_factory=dict)


@dataclass(slots=True)
class _HandlerExecutionContext:
    localizer: "ArtifactLocalizer"
    state: _State
    obj: CarrierObject
    kind: str
    executed: bool = False

    def _execute_once(self) -> None:
        if not self.executed:
            self.executed = True
            self.localizer._dispatch_object(self.obj, self.kind, self.state)

    def extract_records(self, _handler: Any, obj: CarrierObject) -> tuple[CarrierRecord, ...]:
        if obj.id != self.obj.id:
            raise ValueError("handler execution context object mismatch")
        self._execute_once()
        return tuple(
            record for record in self.state.result.records
            if record.metadata.get("owner_object_id") == obj.id
        )

    def enumerate_children(self, _handler: Any, obj: CarrierObject) -> tuple[CarrierObject, ...]:
        if obj.id != self.obj.id:
            raise ValueError("handler execution context object mismatch")
        self._execute_once()
        return tuple(child for child in self.state.objects.values() if child.parent_id == obj.id)


class ArtifactLocalizer:
    def __init__(self, config: AppConfig, registry: HandlerRegistry | None = None):
        self.config = config
        self.registry = registry or default_registry()

    def capabilities(self) -> list[dict[str, Any]]:
        return self.registry.capabilities()

    def scan(
        self, target: Path, timeout_seconds: float | None = None,
        disabled_carrier_types: frozenset[str] = frozenset(),
    ) -> LocalizationResult:
        deadline = time.monotonic() + timeout_seconds if timeout_seconds else None
        state = _State(deadline=deadline, disabled_carrier_types=disabled_carrier_types)
        if target.is_symlink():
            self._warn(state, "symlink_skipped", str(target), "Symbolic links are skipped by default.")
            return state.result
        if target.is_dir():
            for path in sorted(target.rglob("*")):
                if state.total_expanded >= self.config.limits.max_expanded_size:
                    self._warn(state, "expanded_size_limit", str(path), "Total expanded byte limit reached.")
                    break
                try:
                    self._check_timeout(state, str(path))
                except TimeoutError as exc:
                    self._warn(state, "scan_timeout", str(path), str(exc))
                    break
                if path.is_symlink() or not path.is_file():
                    continue
                try:
                    self._read_root(path, state)
                except TimeoutError as exc:
                    self._warn(state, "scan_timeout", str(path), str(exc))
                    break
        elif target.is_file():
            try:
                self._read_root(target, state)
            except TimeoutError as exc:
                self._warn(state, "scan_timeout", str(target), str(exc))
        else:
            self._warn(state, "target_not_found", str(target), "Target does not exist or is not a regular file.")
        return state.result

    def _read_root(self, path: Path, state: _State) -> None:
        try:
            with path.open("rb") as handle:
                data = handle.read(self.config.limits.max_object_size + 1)
        except OSError as exc:
            self._warn(state, "read_error", str(path), str(exc))
            return
        if len(data) > self.config.limits.max_object_size:
            self._warn(state, "object_too_large", str(path), "Root object exceeds max_object_size.")
            return
        self._process_bytes(data, path.name, str(path.resolve()), path.name, 0, state)

    def _process_bytes(
        self,
        data: bytes,
        name: str,
        root_path: str,
        nested_path: str,
        depth: int,
        state: _State,
    ) -> None:
        self._check_timeout(state, nested_path)
        if depth > self.config.limits.max_depth:
            self._warn(state, "max_depth", nested_path, "Nested object depth limit reached.")
            return
        if len(data) > self.config.limits.max_object_size:
            self._warn(state, "object_too_large", nested_path, "Nested object exceeds max_object_size.")
            return
        if state.total_expanded + len(data) > self.config.limits.max_expanded_size:
            self._warn(state, "expanded_size_limit", nested_path, "Total expanded byte limit reached.")
            return
        state.total_expanded += len(data)
        state.result.scanned_objects += 1
        state.object_counter += 1
        content_hash = hashlib.sha256(data).hexdigest()
        parent_nested_path = nested_path.rsplit("!/", 1)[0] if "!/" in nested_path else None
        parent_object_id = state.object_metadata.get(parent_nested_path or "", {}).get("object_id")
        obj = CarrierObject(
            id=f"O{state.object_counter:07d}", root_path=root_path, nested_path=nested_path,
            name=name, data=data, depth=depth, parent_id=parent_object_id, content_sha256=content_hash,
        )
        handler, probe = self.registry.select(obj)
        kind = probe.detected_type
        if kind in state.disabled_carrier_types:
            return
        obj.detected_type = kind
        obj.detected_mime = probe.mime
        structure_valid, structure_evidence = handler.validate_structure(obj, probe)
        state.objects[obj.id] = obj
        state.result.objects.append(obj)
        state.object_metadata[nested_path] = {
            "object_id": obj.id, "parent_object_id": parent_object_id, "mime": probe.mime, "detected_type": kind,
            "content_sha256": content_hash, "handler_id": handler.handler_id,
            "probe_confidence": probe.confidence, "probe_evidence": probe.evidence,
            "structure_valid": structure_valid, "structure_evidence": structure_evidence,
        }
        if not structure_valid:
            self._warn(state, "structure_inconsistent", nested_path, structure_evidence)
        capability = handler.capability()
        if capability.status in {"unavailable", "fallback"} and kind != "binary":
            self._warn(state, "handler_degraded", nested_path, f"{handler.handler_id}: {capability.status}; {capability.reason}")
        execution = _HandlerExecutionContext(self, state, obj, kind)
        tuple(handler.extract_records(obj, execution))
        tuple(handler.enumerate_children(obj, execution))

    def _dispatch_object(self, obj: CarrierObject, kind: str, state: _State) -> None:
        data = obj.data
        root_path = obj.root_path
        nested_path = obj.nested_path
        depth = obj.depth
        handler_started = time.monotonic()
        try:
            if kind in {"text", "json", "yaml", "xml", "csv"}:
                self._handle_text(data, kind, root_path, nested_path, state)
            elif kind in {"zip", "jar", "apk", "tar", "gzip"}:
                self._handle_archive(data, kind, root_path, nested_path, depth, state)
            elif kind in {"docx", "xlsx", "pptx"}:
                self._handle_office(data, kind, root_path, nested_path, depth, state)
            elif kind == "pdf":
                self._handle_pdf(data, root_path, nested_path, depth, state)
            elif kind == "image":
                self._handle_image(data, root_path, nested_path, state, "image")
            elif kind == "sqlite":
                self._handle_sqlite(data, root_path, nested_path, state)
            elif kind == "gettext":
                self._handle_gettext(data, root_path, nested_path, state)
            elif kind in {"7z", "rar", "rpm", "iso"}:
                self._handle_extended_archive(data, kind, root_path, nested_path, depth, state)
            elif kind in {"odf", "epub", "rtf", "outlook-msg", "doc", "xls", "ppt", "ole"}:
                self._handle_document_adapter(data, kind, root_path, nested_path, depth, state)
            elif kind in {"elf", "pe", "mach-o", "wasm", "font", "audio", "video"}:
                self._handle_binary_adapter(data, kind, root_path, nested_path, state)
            else:
                self._handle_binary(data, root_path, nested_path, state)
        except TimeoutError:
            raise
        except Exception as exc:  # Isolate malformed objects from the rest of a scan.
            self._warn(state, "handler_error", nested_path, f"{kind} handler failed: {exc}")
            if kind != "binary":
                try:
                    self._handle_binary(data, root_path, nested_path, state, carrier_type=f"{kind}-fallback")
                except Exception as fallback_exc:
                    self._warn(state, "fallback_error", nested_path, str(fallback_exc))
        finally:
            elapsed = time.monotonic() - handler_started
            if elapsed > self.config.limits.max_handler_seconds:
                self._warn(state, "handler_time_limit", nested_path, f"Handler exceeded {self.config.limits.max_handler_seconds}s budget ({elapsed:.3f}s).")

    def _handle_text(self, data: bytes, kind: str, root_path: str, nested_path: str, state: _State) -> None:
        inference = infer_encoding(data)
        text = inference.text
        encoding_meta = {"encoding": inference.encoding, "encoding_evidence": inference.evidence}
        self._emit(
            state, root_path, nested_path, kind, "document", text, "",
            {"group": nested_path, **encoding_meta}, inference.confidence,
        )
        if kind == "json":
            self._walk_structure(json.loads(text), "$", root_path, nested_path, state)
        elif kind == "yaml":
            for index, document in enumerate(yaml.safe_load_all(text)):
                self._walk_structure(document, f"$[{index}]", root_path, nested_path, state)
        elif kind == "xml":
            root = ElementTree.fromstring(text)
            for index, element in enumerate(root.iter()):
                value = (element.text or "").strip()
                if value:
                    self._emit(
                        state, root_path, nested_path, kind, f"element:{index}", value,
                        element.tag, {"group": nested_path, "field": element.tag, "field_path": element.tag, "relationship": "xml-element"},
                    )
                for key, attr_value in element.attrib.items():
                    self._emit(
                        state, root_path, nested_path, kind, f"attribute:{index}:{key}", attr_value,
                        f"{element.tag}.{key}", {"group": nested_path, "field": key},
                    )
        elif kind == "csv":
            try:
                dialect = csv.Sniffer().sniff(text[:8192])
            except csv.Error:
                dialect = csv.excel
            reader = csv.DictReader(io.StringIO(text), dialect=dialect)
            for row_index, row in enumerate(reader, 2):
                for column, value in row.items():
                    if value:
                        self._emit(
                            state, root_path, nested_path, kind, f"row:{row_index}:column:{column}", value,
                            column or "", {"group": f"{nested_path}:row:{row_index}", "field": column or ""},
                        )
        else:
            for line_number, line in enumerate(text.splitlines(), 1):
                if line.strip():
                    self._emit(
                        state, root_path, nested_path, kind, f"line:{line_number}", line,
                        line[:256], {"group": nested_path, "line": line_number, **encoding_meta},
                    )

    def _walk_structure(
        self,
        value: Any,
        path: str,
        root_path: str,
        nested_path: str,
        state: _State,
    ) -> None:
        if isinstance(value, dict):
            for key, child in value.items():
                self._walk_structure(child, f"{path}.{key}", root_path, nested_path, state)
        elif isinstance(value, list):
            for index, child in enumerate(value):
                self._walk_structure(child, f"{path}[{index}]", root_path, nested_path, state)
        elif value is not None:
            field_name = path.rsplit(".", 1)[-1]
            group = path.rsplit(".", 1)[0] if "." in path else path
            self._emit(
                state, root_path, nested_path, "structured", path, str(value), field_name,
                {"group": f"{nested_path}:{group}", "field": field_name, "field_path": path, "relationship": "field-sibling"},
            )

    def _handle_archive(
        self,
        data: bytes,
        kind: str,
        root_path: str,
        nested_path: str,
        depth: int,
        state: _State,
    ) -> None:
        if kind in {"zip", "jar", "apk"}:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                members = archive.infolist()
                if len(members) > self.config.limits.max_archive_members:
                    self._warn(state, "archive_member_limit", nested_path, "Archive member limit reached.")
                    members = members[: self.config.limits.max_archive_members]
                for member in members:
                    if state.total_expanded >= self.config.limits.max_expanded_size:
                        break
                    if member.is_dir():
                        continue
                    child_path = f"{nested_path}!/{member.filename}"
                    if not self._safe_member_name(member.filename):
                        self._warn(state, "archive_path_traversal", child_path, "Unsafe member path rejected.")
                        continue
                    if member.flag_bits & 0x1:
                        self._warn(state, "encrypted_archive_member", child_path, "Encrypted member skipped.")
                        continue
                    if member.file_size > self.config.limits.max_object_size:
                        self._warn(state, "object_too_large", child_path, "Archive member exceeds max_object_size.")
                        continue
                    compressed = max(1, member.compress_size)
                    if member.file_size / compressed > self.config.limits.max_compression_ratio:
                        self._warn(state, "compression_ratio_limit", child_path, "Archive member compression ratio exceeds limit.")
                        continue
                    with archive.open(member) as handle:
                        child = handle.read(self.config.limits.max_object_size + 1)
                    self._process_bytes(child, member.filename, root_path, child_path, depth + 1, state)
        elif kind == "tar":
            with tarfile.open(fileobj=io.BytesIO(data), mode="r:*") as archive:
                members = archive.getmembers()
                if len(members) > self.config.limits.max_archive_members:
                    self._warn(state, "archive_member_limit", nested_path, "Archive member limit reached.")
                    members = members[: self.config.limits.max_archive_members]
                for member in members:
                    if state.total_expanded >= self.config.limits.max_expanded_size:
                        break
                    if not member.isfile() or member.issym() or member.islnk():
                        continue
                    child_path = f"{nested_path}!/{member.name}"
                    if not self._safe_member_name(member.name):
                        self._warn(state, "archive_path_traversal", child_path, "Unsafe member path rejected.")
                        continue
                    if member.size > self.config.limits.max_object_size:
                        self._warn(state, "object_too_large", child_path, "Archive member exceeds max_object_size.")
                        continue
                    extracted = archive.extractfile(member)
                    if extracted is not None:
                        self._process_bytes(
                            extracted.read(self.config.limits.max_object_size + 1), member.name,
                            root_path, child_path, depth + 1, state,
                        )
        else:
            with gzip.GzipFile(fileobj=io.BytesIO(data)) as archive:
                child = archive.read(self.config.limits.max_object_size + 1)
            if len(child) / max(1, len(data)) > self.config.limits.max_compression_ratio:
                self._warn(state, "compression_ratio_limit", nested_path, "GZIP compression ratio exceeds limit.")
                return
            child_name = Path(nested_path).stem or "gzip-content"
            self._process_bytes(child, child_name, root_path, f"{nested_path}!/{child_name}", depth + 1, state)

    @staticmethod
    def _safe_member_name(name: str) -> bool:
        path = Path(name.replace("\\", "/"))
        return not path.is_absolute() and ".." not in path.parts

    def _handle_extended_archive(
        self, data: bytes, kind: str, root_path: str, nested_path: str, depth: int, state: _State,
    ) -> None:
        """Adapters whose libraries are optional; failure is explicit, then string carving is retained."""
        if kind == "7z":
            import py7zr

            with tempfile.TemporaryDirectory(prefix="divide-7z-") as temp_dir:
                with py7zr.SevenZipFile(io.BytesIO(data), mode="r") as archive:
                    names = archive.getnames()[: self.config.limits.max_archive_members]
                    safe = [name for name in names if self._safe_member_name(name)]
                    for name in set(names) - set(safe):
                        self._warn(state, "archive_path_traversal", f"{nested_path}!/{name}", "Unsafe member path rejected.")
                    archive.extract(path=temp_dir, targets=safe)
                for name in safe:
                    path = Path(temp_dir, name).resolve()
                    base = Path(temp_dir).resolve()
                    if base not in path.parents or not path.is_file():
                        continue
                    if path.stat().st_size > self.config.limits.max_object_size:
                        self._warn(state, "object_too_large", f"{nested_path}!/{name}", "Archive member exceeds limit.")
                        continue
                    self._process_bytes(path.read_bytes(), name, root_path, f"{nested_path}!/{name}", depth + 1, state)
            return
        if kind == "rar":
            import rarfile

            with rarfile.RarFile(io.BytesIO(data)) as archive:
                infos = archive.infolist()[: self.config.limits.max_archive_members]
                for info in infos:
                    name = info.filename
                    child_path = f"{nested_path}!/{name}"
                    if not self._safe_member_name(name):
                        self._warn(state, "archive_path_traversal", child_path, "Unsafe member path rejected.")
                        continue
                    if info.needs_password():
                        self._warn(state, "encrypted_archive_member", child_path, "Encrypted member skipped.")
                        continue
                    if info.file_size > self.config.limits.max_object_size:
                        self._warn(state, "object_too_large", child_path, "Archive member exceeds limit.")
                        continue
                    compressed = max(1, int(getattr(info, "compress_size", info.file_size)))
                    if info.file_size / compressed > self.config.limits.max_compression_ratio:
                        self._warn(state, "compression_ratio_limit", child_path, "RAR member compression ratio exceeds limit.")
                        continue
                    self._process_bytes(archive.read(info), name, root_path, child_path, depth + 1, state)
            return
        if kind == "rpm":
            import rpmfile

            with rpmfile.open(fileobj=io.BytesIO(data)) as archive:
                members = archive.getmembers()[: self.config.limits.max_archive_members]
                for member in members:
                    is_file_attr = getattr(member, "isfile", True)
                    is_file = is_file_attr() if callable(is_file_attr) else bool(is_file_attr)
                    if not is_file or not self._safe_member_name(member.name):
                        continue
                    child = archive.extractfile(member)
                    if child is None:
                        continue
                    payload = child.read(self.config.limits.max_object_size + 1)
                    self._process_bytes(payload, member.name, root_path, f"{nested_path}!/{member.name}", depth + 1, state)
            return
        if kind == "iso":
            import pycdlib

            image = pycdlib.PyCdlib()
            image.open_fp(io.BytesIO(data))
            try:
                count = 0
                for dirname, _dirs, files in image.walk(iso_path="/"):
                    for filename in files:
                        if count >= self.config.limits.max_archive_members:
                            self._warn(state, "archive_member_limit", nested_path, "ISO member limit reached.")
                            return
                        iso_path = f"{dirname.rstrip('/')}/{filename}"
                        output = io.BytesIO()
                        image.get_file_from_iso_fp(output, iso_path=iso_path)
                        payload = output.getvalue()
                        self._process_bytes(payload, str(filename), root_path, f"{nested_path}!/{iso_path.lstrip('/')}", depth + 1, state)
                        count += 1
            finally:
                image.close()

    def _handle_document_adapter(
        self, data: bytes, kind: str, root_path: str, nested_path: str, depth: int, state: _State,
    ) -> None:
        if kind == "odf":
            from odf import teletype
            from odf.opendocument import load
            from odf.text import P

            document = load(io.BytesIO(data))
            for index, node in enumerate(document.getElementsByType(P)):
                text = teletype.extractText(node)
                if text.strip():
                    self._emit(state, root_path, nested_path, kind, f"paragraph:{index}", text, "ODF paragraph", {"group": nested_path})
            self._process_office_embeds(data, root_path, nested_path, depth, state)
            return
        if kind == "epub":
            from bs4 import BeautifulSoup
            from ebooklib import ITEM_DOCUMENT, epub

            book = epub.read_epub(io.BytesIO(data), options={"ignore_ncx": True})
            for index, item in enumerate(book.get_items_of_type(ITEM_DOCUMENT)):
                text = BeautifulSoup(item.get_content(), "html.parser").get_text(" ", strip=True)
                if text:
                    self._emit(state, root_path, nested_path, kind, f"document:{index}", text, item.get_name(), {"group": nested_path})
            self._handle_archive(data, "zip", root_path, nested_path, depth, state)
            return
        if kind == "rtf":
            from striprtf.striprtf import rtf_to_text

            text = rtf_to_text(data.decode("latin-1"))
            self._emit(state, root_path, nested_path, kind, "document", text, "RTF document", {"group": nested_path})
            return
        if kind == "xls":
            import xlrd

            workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
            for sheet in workbook.sheets():
                for row_index in range(min(sheet.nrows, self.config.limits.max_sqlite_rows)):
                    for column_index, value in enumerate(sheet.row_values(row_index)):
                        if value not in (None, ""):
                            self._emit(
                                state, root_path, nested_path, kind,
                                f"sheet:{sheet.name}:row:{row_index}:column:{column_index}", str(value), sheet.name,
                                {"group": f"{nested_path}:sheet:{sheet.name}:row:{row_index}", "field": str(column_index)},
                            )
            return
        if kind == "outlook-msg":
            import extract_msg

            with tempfile.NamedTemporaryFile(prefix="divide-msg-", suffix=".msg", delete=False) as handle:
                handle.write(data)
                temp_path = handle.name
            try:
                message = extract_msg.Message(temp_path)
                fields = {"subject": message.subject, "sender": message.sender, "to": message.to, "cc": message.cc, "body": message.body}
                for key, value in fields.items():
                    if value:
                        self._emit(state, root_path, nested_path, kind, f"field:{key}", str(value), key, {"group": nested_path, "field": key})
                for index, attachment in enumerate(message.attachments):
                    payload = getattr(attachment, "data", None)
                    name = getattr(attachment, "longFilename", None) or getattr(attachment, "shortFilename", None) or f"attachment-{index}"
                    if isinstance(payload, bytes):
                        self._process_bytes(payload, name, root_path, f"{nested_path}!/{name}", depth + 1, state)
                message.close()
            finally:
                try:
                    os.unlink(temp_path)
                except OSError:
                    pass
            return
        # DOC/PPT/unknown OLE: expose stream names, then use bounded byte-string fallback.
        try:
            import olefile

            with olefile.OleFileIO(io.BytesIO(data)) as document:
                for index, stream in enumerate(document.listdir(streams=True, storages=False)):
                    self._emit(
                        state, root_path, nested_path, kind, f"ole-stream:{index}", "/".join(stream),
                        "OLE stream name", {"group": nested_path, "relationship": "stream"}, .65,
                    )
        finally:
            self._handle_binary(data, root_path, nested_path, state, carrier_type=kind)

    def _handle_binary_adapter(
        self, data: bytes, kind: str, root_path: str, nested_path: str, state: _State,
    ) -> None:
        if kind == "font":
            from fontTools.ttLib import TTFont

            font = TTFont(io.BytesIO(data), lazy=True)
            if "name" in font:
                for index, entry in enumerate(font["name"].names):
                    try:
                        value = entry.toUnicode()
                    except UnicodeError:
                        continue
                    self._emit(state, root_path, nested_path, kind, f"name:{index}", value, "font name table", {"group": nested_path}, .8)
            font.close()
        elif kind == "audio":
            import mutagen

            media = mutagen.File(io.BytesIO(data), easy=True)
            for key, values in (media or {}).items():
                for index, value in enumerate(values if isinstance(values, list) else [values]):
                    self._emit(state, root_path, nested_path, kind, f"tag:{key}:{index}", str(value), key, {"group": nested_path, "field": key}, .85)
        elif kind == "video":
            try:
                from pymediainfo import MediaInfo

                info = MediaInfo.parse(io.BytesIO(data), output="JSON")
                self._emit(state, root_path, nested_path, kind, "metadata", str(info), "media metadata", {"group": nested_path}, .65)
            except (ImportError, OSError, RuntimeError) as exc:
                self._warn(state, "video_metadata_fallback", nested_path, str(exc))
        else:
            header = {"format": kind, "size": len(data), "magic_hex": data[:16].hex()}
            try:
                import lief

                with tempfile.NamedTemporaryFile(prefix="divide-executable-", delete=False) as handle:
                    handle.write(data)
                    parse_path = handle.name
                try:
                    executable = lief.parse(parse_path)
                    if executable is not None:
                        header["entrypoint"] = int(getattr(executable, "entrypoint", 0))
                        header["sections"] = [section.name for section in getattr(executable, "sections", [])][:256]
                finally:
                    try:
                        os.unlink(parse_path)
                    except OSError:
                        pass
            except (ImportError, OSError, ValueError, RuntimeError) as exc:
                header["parser_fallback"] = str(exc)
            self._emit(state, root_path, nested_path, kind, "binary-header", json.dumps(header), "binary metadata", {"group": nested_path}, .7)
        self._handle_binary(data, root_path, nested_path, state, carrier_type=kind)

    def _handle_office(
        self,
        data: bytes,
        kind: str,
        root_path: str,
        nested_path: str,
        depth: int,
        state: _State,
    ) -> None:
        if kind == "docx":
            from docx import Document

            document = Document(io.BytesIO(data))
            for index, paragraph in enumerate(document.paragraphs):
                if paragraph.text.strip():
                    self._emit(state, root_path, nested_path, kind, f"paragraph:{index}", paragraph.text, "paragraph", {"group": nested_path})
            for table_index, table in enumerate(document.tables):
                for row_index, row in enumerate(table.rows):
                    for column_index, cell in enumerate(row.cells):
                        if cell.text.strip():
                            self._emit(
                                state, root_path, nested_path, kind,
                                f"table:{table_index}:row:{row_index}:column:{column_index}", cell.text,
                                f"table {table_index}", {"group": f"{nested_path}:table:{table_index}:row:{row_index}"},
                            )
        elif kind == "xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            self._emit(
                                state, root_path, nested_path, kind, f"sheet:{sheet.title}:cell:{cell.coordinate}",
                                str(cell.value), cell.coordinate,
                                {"group": f"{nested_path}:sheet:{sheet.title}:row:{cell.row}", "field": cell.coordinate},
                            )
        else:
            from pptx import Presentation

            presentation = Presentation(io.BytesIO(data))
            for slide_index, slide in enumerate(presentation.slides, 1):
                for shape_index, shape in enumerate(slide.shapes):
                    if hasattr(shape, "text") and shape.text.strip():
                        self._emit(
                            state, root_path, nested_path, kind, f"slide:{slide_index}:shape:{shape_index}",
                            shape.text, f"slide {slide_index}", {"group": f"{nested_path}:slide:{slide_index}"},
                        )
        self._process_office_embeds(data, root_path, nested_path, depth, state)

    def _process_office_embeds(
        self, data: bytes, root_path: str, nested_path: str, depth: int, state: _State,
    ) -> None:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            for member in archive.infolist():
                if state.total_expanded >= self.config.limits.max_expanded_size:
                    break
                normalized = member.filename.lower()
                if not any(part in normalized for part in ("/media/", "/embeddings/", "pictures/", "objectreplacements/")) or member.is_dir():
                    continue
                child_path = f"{nested_path}!/{member.filename}"
                if not self._safe_member_name(member.filename):
                    self._warn(state, "archive_path_traversal", child_path, "Unsafe embedded-object path rejected.")
                    continue
                if member.file_size > self.config.limits.max_object_size:
                    self._warn(state, "object_too_large", child_path, "Embedded Office object exceeds limit.")
                    continue
                if member.file_size / max(1, member.compress_size) > self.config.limits.max_compression_ratio:
                    self._warn(state, "compression_ratio_limit", child_path, "Embedded object compression ratio exceeds limit.")
                    continue
                self._process_bytes(archive.read(member), member.filename, root_path, child_path, depth + 1, state)

    def _handle_pdf(
        self, data: bytes, root_path: str, nested_path: str, depth: int, state: _State,
    ) -> None:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            self._warn(state, "encrypted_pdf", nested_path, "Encrypted PDF skipped.")
            return
        page_text: dict[int, str] = {}
        for page_index, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            page_text[page_index] = text
            if text.strip():
                self._emit(
                    state, root_path, nested_path, "pdf", f"page:{page_index}", text,
                    f"PDF page {page_index}", {"group": f"{nested_path}:page:{page_index}", "page": page_index, "relationship": "page-text"},
                )
            for image_index, embedded in enumerate(page.images):
                image_name = embedded.name or f"page-{page_index}-image-{image_index}.bin"
                self._process_bytes(
                    embedded.data, image_name, root_path,
                    f"{nested_path}!/page-{page_index}/{image_name}", depth + 1, state,
                )
        if self.config.ocr.enabled:
            import fitz

            document = fitz.open(stream=data, filetype="pdf")
            for page_index, page in enumerate(document, 1):
                if page_text.get(page_index, "").strip() or state.ocr_pages >= self.config.limits.max_ocr_pages:
                    continue
                pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                self._handle_image(
                    pixmap.tobytes("png"), root_path, f"{nested_path}!/rendered-page-{page_index}.png",
                    state, "pdf-page",
                )

    def _handle_image(
        self, data: bytes, root_path: str, nested_path: str, state: _State, carrier_type: str,
    ) -> None:
        if not self.config.ocr.enabled:
            return
        if state.ocr_pages >= self.config.limits.max_ocr_pages:
            self._warn(state, "ocr_page_limit", nested_path, "OCR page/image limit reached.")
            return
        from PIL import Image
        import pytesseract
        from pytesseract import Output

        state.ocr_pages += 1
        source_image = Image.open(io.BytesIO(data))
        if source_image.width * source_image.height > self.config.limits.max_image_pixels:
            self._warn(state, "image_pixel_limit", nested_path, "Image pixel budget exceeded; OCR skipped.")
            return
        frame_count = int(getattr(source_image, "n_frames", 1))
        if frame_count > self.config.limits.max_media_frames:
            self._warn(
                state, "media_frame_limit", nested_path,
                f"Image has {frame_count} frames; only the first frame is processed under the media-frame budget.",
            )
        image = source_image.convert("RGB")
        result = pytesseract.image_to_data(image, lang=self.config.ocr.language, output_type=Output.DICT)
        lines: dict[tuple[int, int, int], list[tuple[int, str, float, dict[str, int]]]] = defaultdict(list)
        for index, raw_text in enumerate(result.get("text", [])):
            word = raw_text.strip()
            if not word:
                continue
            try:
                confidence = float(result["conf"][index])
            except (ValueError, TypeError):
                confidence = 0.0
            if confidence < self.config.ocr.min_confidence:
                continue
            key = (result["block_num"][index], result["par_num"][index], result["line_num"][index])
            box = {name: int(result[name][index]) for name in ("left", "top", "width", "height")}
            lines[key].append((box["left"], word, confidence, box))
        for line_index, key in enumerate(sorted(lines), 1):
            words = sorted(lines[key], key=lambda item: item[0])
            text, line_bbox = reconstruct_ocr_line(words)
            average = sum(item[2] for item in words) / len(words) / 100.0
            self._emit(
                state, root_path, nested_path, carrier_type, f"ocr-line:{line_index}", text,
                "OCR", {
                    "group": f"{nested_path}:block:{key[0]}:paragraph:{key[1]}",
                    "ocr_boxes": [item[3] for item in words], "ocr_bbox": line_bbox,
                    "ocr_block": key[0], "ocr_paragraph": key[1], "ocr_line": key[2],
                    "relationship": "ocr-line",
                }, average,
            )

    def _handle_sqlite(self, data: bytes, root_path: str, nested_path: str, state: _State) -> None:
        temp_path = ""
        try:
            with tempfile.NamedTemporaryFile(prefix="divide-sqlite-", suffix=".db", delete=False) as handle:
                handle.write(data)
                temp_path = handle.name
            uri = f"file:{Path(temp_path).as_posix()}?mode=ro&immutable=1"
            connection = sqlite3.connect(uri, uri=True)
            connection.execute("PRAGMA query_only=ON")
            vm_steps = 0

            def budget_guard() -> int:
                nonlocal vm_steps
                vm_steps += 1_000
                timed_out = bool(state.deadline and time.monotonic() > state.deadline)
                return int(timed_out or vm_steps > self.config.limits.max_sqlite_vm_steps)

            connection.set_progress_handler(budget_guard, 1_000)
            tables = connection.execute(
                "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'"
            ).fetchall()
            remaining = self.config.limits.max_sqlite_rows
            for (table_name,) in tables:
                if remaining <= 0:
                    break
                escaped = str(table_name).replace('"', '""')
                cursor = connection.execute(f'SELECT * FROM "{escaped}" LIMIT ?', (remaining,))
                columns = [column[0] for column in cursor.description or []]
                for row_index, row in enumerate(cursor, 1):
                    remaining -= 1
                    for column, value in zip(columns, row, strict=False):
                        if value is not None:
                            self._emit(
                                state, root_path, nested_path, "sqlite",
                                f"table:{table_name}:row:{row_index}:column:{column}", str(value), column,
                                {"group": f"{nested_path}:table:{table_name}:row:{row_index}", "field": column},
                            )
            connection.close()
            if remaining <= 0:
                self._warn(state, "sqlite_row_limit", nested_path, "SQLite row limit reached.")
        finally:
            if temp_path:
                try:
                    os.unlink(temp_path)
                except OSError:
                    pass

    def _handle_gettext(self, data: bytes, root_path: str, nested_path: str, state: _State) -> None:
        translations = gettext.GNUTranslations(io.BytesIO(data))
        catalog = getattr(translations, "_catalog", {})
        for index, (source, translated) in enumerate(catalog.items()):
            if not isinstance(source, str):
                continue
            values: Iterable[Any] = translated if isinstance(translated, tuple) else (translated,)
            for value_index, value in enumerate(values):
                self._emit(
                    state, root_path, nested_path, "gettext", f"message:{index}:{value_index}",
                    str(value), source, {"group": nested_path, "field": source},
                )

    def _handle_binary(
        self, data: bytes, root_path: str, nested_path: str, state: _State, carrier_type: str = "binary",
    ) -> None:
        for index, match in enumerate(re.finditer(rb"[\x20-\x7e]{6,}", data)):
            self._emit(
                state, root_path, nested_path, carrier_type, f"ascii-offset:{match.start()}",
                match.group().decode("ascii", errors="ignore"), "binary string",
                {"group": nested_path, "offset": match.start(), "string_index": index}, 0.8,
            )
        for index, match in enumerate(re.finditer(rb"(?:[\x20-\x7e]\x00){6,}", data)):
            self._emit(
                state, root_path, nested_path, carrier_type, f"utf16-offset:{match.start()}",
                match.group().decode("utf-16le", errors="ignore"), "binary UTF-16 string",
                {"group": nested_path, "offset": match.start(), "string_index": index}, 0.75,
            )
        for index, match in enumerate(re.finditer(rb"(?:\x00[\x20-\x7e]){6,}", data)):
            self._emit(
                state, root_path, nested_path, carrier_type, f"utf16be-offset:{match.start()}",
                match.group().decode("utf-16be", errors="ignore"), "binary UTF-16BE string",
                {"group": nested_path, "offset": match.start(), "string_index": index}, .75,
            )

    def _emit(
        self,
        state: _State,
        root_path: str,
        nested_path: str,
        carrier_type: str,
        location: str,
        text: str,
        context: str,
        metadata: dict[str, Any],
        confidence: float = 1.0,
    ) -> None:
        if not text:
            return
        state.record_counter += 1
        object_meta = state.object_metadata.get(nested_path, {})
        metadata = {
            **metadata, "owner_object_id": object_meta.get("object_id"),
            "handler_id": object_meta.get("handler_id"), "probe_evidence": object_meta.get("probe_evidence", []),
            "structure_valid": object_meta.get("structure_valid"),
            "structure_evidence": object_meta.get("structure_evidence"),
        }
        encoding = metadata.get("encoding")
        field_path = metadata.get("field_path", metadata.get("field"))
        bbox = metadata.get("ocr_bbox")
        page = metadata.get("page")
        if page is None:
            page_match = re.search(r"(?:rendered-)?page-(\d+)", nested_path)
            page = int(page_match.group(1)) if page_match else None
        text_hash = hashlib.sha256(text.encode("utf-8", errors="replace")).hexdigest()
        state.result.records.append(
            CarrierRecord(
                id=f"R{state.record_counter:07d}", root_path=root_path, nested_path=nested_path,
                carrier_type=carrier_type, location=location, text=text, context=context,
                confidence=max(0.0, min(1.0, confidence)), metadata=metadata,
                parent_object_id=object_meta.get("parent_object_id"), mime=object_meta.get("mime"),
                detected_type=object_meta.get("detected_type", carrier_type), encoding=encoding,
                byte_offset=metadata.get("offset"), page=page, field_path=field_path,
                ocr_bbox=tuple(bbox) if bbox else None, relationship=metadata.get("relationship"),
                content_sha256=text_hash, extraction_confidence=max(0.0, min(1.0, confidence)),
                encoding_evidence=metadata.get("encoding_evidence", {}),
            )
        )

    @staticmethod
    def _warn(state: _State, code: str, path: str, message: str) -> None:
        state.result.warnings.append(ScanWarning(code=code, path=path, message=message))

    @staticmethod
    def _check_timeout(state: _State, path: str) -> None:
        if state.deadline and time.monotonic() > state.deadline:
            raise TimeoutError(f"Scan timeout reached while processing {path}")

